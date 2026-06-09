import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Brand Colors
    BG_DARK = RGBColor(26, 20, 40)        # #1A1428 (Dark Purple)
    CARD_BG = RGBColor(42, 31, 61)        # #2A1F3D (Darker Card)
    PRIMARY_PURPLE = RGBColor(123, 91, 168) # #7B5BA8 (Primary Purple)
    PRIMARY_LIGHT = RGBColor(155, 127, 196) # #9B7FC4 (Light Purple)
    ACCENT_YELLOW = RGBColor(244, 211, 94) # #F4D35E (Accent Yellow)
    WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(176, 168, 195)  # #B0A8C3 (Muted Lavender)
    SUCCESS_GREEN = RGBColor(46, 204, 113) # #2ECC71
    ERROR_RED = RGBColor(231, 76, 60)      # #E74C3C
    
    # Helper to set slide background
    def set_dark_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    # Helper to create a slide with standard header
    def add_slide_with_header(title_text):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        set_dark_bg(slide)
        
        # Add Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Montserrat'
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = ACCENT_YELLOW
        
        # Add thin line below header
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = PRIMARY_PURPLE
        line.line.fill.background()
        
        return slide

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    slide_layout = prs.slide_layouts[6] # blank layout
    slide1 = prs.slides.add_slide(slide_layout)
    set_dark_bg(slide1)
    
    # Large Decorative Background Card
    bg_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    bg_card.fill.solid()
    bg_card.fill.fore_color.rgb = CARD_BG
    bg_card.line.color.rgb = PRIMARY_PURPLE
    bg_card.line.width = Pt(2)
    
    # App Icon Placeholder (Shape)
    icon_shape = slide1.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(5.666), Inches(2.0), Inches(2.0), Inches(1.8))
    icon_shape.fill.solid()
    icon_shape.fill.fore_color.rgb = ACCENT_YELLOW
    icon_shape.line.fill.background()
    tf_icon = icon_shape.text_frame
    tf_icon.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_icon = tf_icon.paragraphs[0]
    p_icon.text = "Rp"
    p_icon.alignment = PP_ALIGN.CENTER
    p_icon.font.name = 'Montserrat'
    p_icon.font.size = Pt(40)
    p_icon.font.bold = True
    p_icon.font.color.rgb = BG_DARK
    
    # Title Text Box
    title_box = slide1.shapes.add_textbox(Inches(1.5), Inches(4.0), Inches(10.333), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FINANCE TRACKER"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Montserrat'
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = "Solusi Digital Perhitungan HPP, BEP, dan Manajemen Keuangan UMKM"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = 'Inter'
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_YELLOW
    
    # Subtitle / Presenter info
    info_box = slide1.shapes.add_textbox(Inches(1.5), Inches(5.3), Inches(10.333), Inches(0.5))
    tf_info = info_box.text_frame
    p_info = tf_info.paragraphs[0]
    p_info.text = "Pengembangan Sistem Informasi Bisnis Mandiri berbasis Mobile Flutter"
    p_info.alignment = PP_ALIGN.CENTER
    p_info.font.name = 'Inter'
    p_info.font.size = Pt(14)
    p_info.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: LATAR BELAKANG & MASALAH UMKM
    # ==========================================
    slide2 = add_slide_with_header("Latar Belakang & Permasalahan UMKM")
    
    # Subtitle text
    sub_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.4))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = "Pelaku UMKM sering mengalami kegagalan modal dan profit akibat 4 pilar masalah manajemen keuangan:"
    sub_p.font.name = 'Inter'
    sub_p.font.size = Pt(16)
    sub_p.font.color.rgb = TEXT_MUTED
    
    # 4 Cards for 4 Problems
    problems = [
        ("HPP Tidak Akurat", "Menetapkan harga jual secara tebak-tebakan tanpa menghitung biaya bahan baku riil, upah tenaga kerja, dan biaya overhead pendukung secara presisi."),
        ("Kebutaan Titik Impas (BEP)", "Menjalankan usaha tanpa mengetahui batas minimum volume penjualan produk yang harus dicapai agar tidak mengalami kerugian operasional."),
        ("Tabungan & Modal Bercampur", "Uang pribadi dan keuntungan usaha menyatu di rekening yang sama. Sulit menyisihkan laba secara konsisten untuk rencana ekspansi bisnis."),
        ("Pencatatan Konvensional", "Mengandalkan pencatatan kertas fisik yang rentan hilang, rusak, salah hitung, serta tidak memiliki visualisasi statistik tren performa finansial.")
    ]
    
    for i, (title, desc) in enumerate(problems):
        row = i // 2
        col = i % 2
        
        left = Inches(0.8 + col * 5.95)
        top = Inches(2.0 + row * 2.5)
        width = Inches(5.6)
        height = Inches(2.2)
        
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.3)
        tf_card.margin_top = Inches(0.25)
        tf_card.margin_right = Inches(0.3)
        tf_card.margin_bottom = Inches(0.2)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = f"{i+1}. {title}"
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(10)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 3: SOLUSI YANG DITAWARKAN
    # ==========================================
    slide3 = add_slide_with_header("Solusi Aplikasi - Finance Tracker")
    
    # Left Column: Highlight Text
    left_box = slide3.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(4.5), Inches(4.5))
    ltf = left_box.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "Transformasi Keuangan UMKM Secara Mandiri"
    lp1.font.name = 'Montserrat'
    lp1.font.size = Pt(28)
    lp1.font.bold = True
    lp1.font.color.rgb = WHITE
    lp1.space_after = Pt(20)
    
    lp2 = ltf.add_paragraph()
    lp2.text = "Finance Tracker menyediakan ekosistem kalkulator keuangan, perencanaan alokasi tabungan, dan visualisasi performa bisnis dalam genggaman tangan."
    lp2.font.name = 'Inter'
    lp2.font.size = Pt(16)
    lp2.font.color.rgb = TEXT_MUTED
    
    # Right Column: 4 key solution points in cards
    solutions = [
        ("Automasi HPP & BEP", "Menghitung HPP unit/total serta BEP unit/omzet secara real-time dan akurat."),
        ("Simulasi Slider Dinamis", "Mensimulasikan perubahan harga jual untuk memantau fluktuasi margin laba & BEP."),
        ("Savings Plan Terproteksi", "Menyisihkan laba ke wadah tabungan khusus bisnis dengan alokasi target dan tenggat."),
        ("Insight Grafik Interaktif", "Visualisasi tren HPP, BEP, dan performa keuangan untuk keputusan bisnis cerdas.")
    ]
    
    for i, (title, desc) in enumerate(solutions):
        left = Inches(5.8)
        top = Inches(1.8 + i * 1.3)
        width = Inches(6.7)
        height = Inches(1.1)
        
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1)
        
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.2)
        tf_card.margin_top = Inches(0.15)
        tf_card.margin_right = Inches(0.2)
        tf_card.margin_bottom = Inches(0.1)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 4: ALUR BISNIS SISTEM
    # ==========================================
    slide4 = add_slide_with_header("Alur Bisnis Keuangan UMKM (Business Flow)")
    
    # Subtitle
    sub_box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.4))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = "Prosedur operasional keuangan harian dan bulanan yang diakomodasi oleh sistem:"
    sub_p.font.name = 'Inter'
    sub_p.font.size = Pt(15)
    sub_p.font.color.rgb = TEXT_MUTED
    
    # 5 Horizontal Steps
    steps = [
        ("1. Input Biaya", "Input bahan baku, upah tenaga kerja, overhead & biaya tetap bulanan."),
        ("2. Simulasi Harga", "Gunakan slider harga jual untuk simulasi keuntungan & margin."),
        ("3. Simpan Riwayat", "Simpan data perhitungan HPP & BEP ke cloud history secara aman."),
        ("4. Tabung Laba", "Alokasikan keuntungan bersih ke rencana tabungan bisnis terpisah."),
        ("5. Evaluasi Insight", "Pantau grafik tren bulanan & ekspor laporan PDF sebagai evaluasi.")
    ]
    
    for i, (title, desc) in enumerate(steps):
        left = Inches(0.8 + i * 2.38)
        top = Inches(2.5)
        width = Inches(2.2)
        height = Inches(3.8)
        
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        # Highlight first and last with yellow accent top border
        if i == 0 or i == 4:
            accent_bar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.15))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = ACCENT_YELLOW
            accent_bar.line.fill.background()
            
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.15)
        tf_card.margin_top = Inches(0.35)
        tf_card.margin_right = Inches(0.15)
        tf_card.margin_bottom = Inches(0.15)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(12)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = WHITE
        
        # Draw Arrow connector between steps (except last one)
        if i < 4:
            arrow = slide4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left + Inches(2.2) + Inches(0.04), top + Inches(1.8), Inches(0.1), Inches(0.2))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = PRIMARY_PURPLE
            arrow.line.fill.background()

    # ==========================================
    # SLIDE 5: FITUR UTAMA APLIKASI
    # ==========================================
    slide5 = add_slide_with_header("Fitur Utama Aplikasi")
    
    features = [
        ("Dashboard Finansial", "Memantau target profit bulanan vs untung bersih riil saat ini. Dilengkapi progress bar visual yang interaktif serta dukungan seeder data sampel untuk pengguna baru."),
        ("Kalkulator HPP & BEP", "Kalkulator terstruktur 2 langkah. Langkah 1 menghitung biaya produksi untuk HPP unit. Langkah 2 mengombinasikan biaya tetap & harga jual untuk menghitung BEP Unit & Omzet."),
        ("Simulasi Harga Jual", "Slider harga jual yang interaktif. Memberikan perubahan profit margin secara instan dan badge kategori dinamis: RUGI, KEUNTUNGAN RENDAH, SEDANG, atau TINGGI."),
        ("Rencana Tabungan (Savings)", "Sistem pencatatan tabungan bisnis mandiri. Fitur deposit (tabung) dan withdrawal (tarik) saldo tabungan, dengan tracking persentase progres pencapaian."),
        ("Grafik Insight Finansial", "Menyajikan grafik komparasi HPP vs BEP bulanan, laju pertumbuhan akumulasi seluruh tabungan, analisis efisiensi biaya produksi, dan statistik produk terpopuler."),
        ("Premium Membership 👑", "Skema upgrade akun. Pengguna gratis dibatasi maksimal 3 data history. Pengguna premium mendapatkan penyimpanan tak terbatas, bebas batasan, dan fitur ekspor PDF.")
    ]
    
    for i, (title, desc) in enumerate(features):
        row = i // 3
        col = i % 3
        
        left = Inches(0.8 + col * 3.95)
        top = Inches(1.8 + row * 2.6)
        width = Inches(3.75)
        height = Inches(2.3)
        
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        # Gold badge for Premium
        if "Premium" in title:
            card.line.color.rgb = ACCENT_YELLOW
            badge = slide5.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(3.2), top - Inches(0.1), Inches(0.6), Inches(0.4))
            badge.fill.solid()
            badge.fill.fore_color.rgb = ACCENT_YELLOW
            badge.line.fill.background()
            btf = badge.text_frame
            bp = btf.paragraphs[0]
            bp.text = "PRO"
            bp.font.name = 'Montserrat'
            bp.font.size = Pt(10)
            bp.font.bold = True
            bp.font.color.rgb = BG_DARK
            bp.alignment = PP_ALIGN.CENTER
            
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.2)
        tf_card.margin_top = Inches(0.2)
        tf_card.margin_right = Inches(0.2)
        tf_card.margin_bottom = Inches(0.15)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(10)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 6: ANALISIS SISTEM - USE CASE DIAGRAM
    # ==========================================
    slide6 = add_slide_with_header("Analisis Sistem: Use Case Diagram")
    
    # Left text description
    left_box = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.5), Inches(4.8))
    ltf = left_box.text_frame
    ltf.word_wrap = True
    
    p = ltf.paragraphs[0]
    p.text = "Interaksi Aktor & Fungsionalitas"
    p.font.name = 'Montserrat'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(15)
    
    bullet_points = [
        "Aktor Sistem: Pelaku UMKM terbagi menjadi Free Member dan Premium Member.",
        "Free Member memiliki akses ke semua modul inti (kalkulator, tabungan, insight), namun dibatasi hanya dapat menyimpan maksimal 3 riwayat perhitungan HPP & BEP.",
        "Premium Member memiliki akses tanpa batas (unlimited history) dan hak eksklusif untuk mengekspor laporan keuangan dalam format PDF.",
        "Aktor Pendukung: Firebase Authentication mengamankan akses akun & Firestore mengelola database tersinkronisasi."
    ]
    
    for bp in bullet_points:
        p_bp = ltf.add_paragraph()
        p_bp.text = "• " + bp
        p_bp.font.name = 'Inter'
        p_bp.font.size = Pt(12)
        p_bp.font.color.rgb = TEXT_MUTED
        p_bp.space_after = Pt(8)
        
    # Right Box simulating Use Case Layout Visual
    right_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.8), Inches(1.8), Inches(6.7), Inches(4.8))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = CARD_BG
    right_card.line.color.rgb = PRIMARY_PURPLE
    right_card.line.width = Pt(1.5)
    
    rtf = right_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.3)
    rtf.margin_top = Inches(0.25)
    rtf.margin_right = Inches(0.3)
    
    rp = rtf.paragraphs[0]
    rp.text = "Peta Kasus Penggunaan (Use Case Mapping)"
    rp.font.name = 'Montserrat'
    rp.font.size = Pt(16)
    rp.font.bold = True
    rp.font.color.rgb = ACCENT_YELLOW
    rp.space_after = Pt(15)
    
    uc_items = [
        ("Autentikasi Pengguna", "Registrasi, Login, Google Sign-In, Verifikasi Email, Reset Kata Sandi."),
        ("Kalkulasi & Simulasi", "Formulir HPP (biaya produksi) & BEP (biaya tetap, harga unit) serta slider simulasi margin."),
        ("Manajemen Tabungan", "Membuat rencana tabungan dengan target dana, tanggal jatuh tempo, deposit & withdraw saldo."),
        ("Dashboard & Insight", "Mengubah target profit bulanan, melihat progress bar, dan statistik visual grafik interaktif."),
        ("Fitur Khusus Premium", "Upgrade keanggotaan (simulasi bayar) & ekspor laporan PDF (PdfService).")
    ]
    
    for item_title, item_desc in uc_items:
        p_item = rtf.add_paragraph()
        p_item.text = f"✔ {item_title}: "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = WHITE
        
        p_desc = rtf.add_paragraph()
        p_desc.text = item_desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(10)

    # ==========================================
    # SLIDE 7: ACTIVITY DIAGRAM - HPP & BEP
    # ==========================================
    slide7 = add_slide_with_header("Perancangan: Workflow Perhitungan HPP & BEP")
    
    sub_box = slide7.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.4))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = "Alur aktivitas pengguna saat memasukkan biaya produksi hingga pengecekan kuota riwayat:"
    sub_p.font.name = 'Inter'
    sub_p.font.size = Pt(15)
    sub_p.font.color.rgb = TEXT_MUTED
    
    # 4 Workflow Steps vertically stacked or horizontally
    stages = [
        ("Langkah 1: Input Data HPP", "Pengguna mengisi form biaya bahan baku, upah tenaga kerja, overhead pabrik, dan jumlah unit produksi. Sistem menghitung HPP per Unit."),
        ("Langkah 2: Input Data BEP & Jual", "Pengguna mengisi form biaya tetap operasional usaha dan harga jual unit. Sistem menghitung titik impas BEP Unit dan BEP Rupiah."),
        ("Langkah 3: Simulasi Slider", "Pengguna menggeser slider harga jual. Sistem menghitung ulang margin laba secara real-time dan memperbarui label margin (Rugi s/d Tinggi)."),
        ("Langkah 4: Validasi Kuota Simpan", "Sistem memvalidasi status user. Jika pengguna Gratis dan data history >= 3, simpan diblokir dan diarahkan untuk upgrade Premium. Jika Premium, data tersimpan.")
    ]
    
    for i, (title, desc) in enumerate(stages):
        left = Inches(0.8 + (i % 2) * 5.95)
        top = Inches(1.9 + (i // 2) * 2.5)
        width = Inches(5.6)
        height = Inches(2.2)
        
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        if "Validasi" in title:
            card.line.color.rgb = ACCENT_YELLOW
            
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.25)
        tf_card.margin_top = Inches(0.2)
        tf_card.margin_right = Inches(0.25)
        tf_card.margin_bottom = Inches(0.15)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(10)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 8: ACTIVITY DIAGRAM - SAVINGS PLAN
    # ==========================================
    slide8 = add_slide_with_header("Perancangan: Workflow Savings Plan (Tabungan)")
    
    sub_box = slide8.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.4))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = "Alur kerja pengelolaan tabungan bisnis, dari inisiasi target hingga deposit/penarikan dana:"
    sub_p.font.name = 'Inter'
    sub_p.font.size = Pt(15)
    sub_p.font.color.rgb = TEXT_MUTED
    
    stages_saving = [
        ("Langkah 1: Pembuatan Rencana", "Pengguna menginput rencana tabungan (nama, target dana, saldo awal, tenggat waktu) di menu Wallet. Data tersimpan di Firestore."),
        ("Langkah 2: Tampilan Progress", "Sistem memuat daftar rencana tabungan, menghitung sisa dana yang dibutuhkan, dan persentase progress pencapaian target tabungan secara otomatis."),
        ("Langkah 3: Alur Deposit (Tabung)", "Pengguna menambahkan nominal tabungan. Sistem memperbarui data 'currentAmount' dan menambahkan log transaksi riil tabungan."),
        ("Langkah 4: Alur Withdrawal (Tarik)", "Pengguna menarik dana jika darurat. Sistem memvalidasi saldo yang cukup, memotong 'currentAmount', dan memperbarui tampilan progres.")
    ]
    
    for i, (title, desc) in enumerate(stages_saving):
        left = Inches(0.8 + (i % 2) * 5.95)
        top = Inches(1.9 + (i // 2) * 2.5)
        width = Inches(5.6)
        height = Inches(2.2)
        
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.25)
        tf_card.margin_top = Inches(0.2)
        tf_card.margin_right = Inches(0.25)
        tf_card.margin_bottom = Inches(0.15)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(10)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 9: PERANCANGAN - CLASS DIAGRAM (MVVM)
    # ==========================================
    slide9 = add_slide_with_header("Perancangan: Class Diagram (Arsitektur MVVM)")
    
    sub_box = slide9.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.4))
    sub_tf = sub_box.text_frame
    sub_p = sub_tf.paragraphs[0]
    sub_p.text = "Pemetaan kelas di codebase Flutter terbagi menjadi Model, Service (Data), dan ViewModel (Logic):"
    sub_p.font.name = 'Inter'
    sub_p.font.size = Pt(15)
    sub_p.font.color.rgb = TEXT_MUTED
    
    # 3 Large Vertical Panels representing MVVM layers
    layers = [
        ("1. MODEL LAYER (Data Entities)", 
         "Menampung model data yang dipetakan dari Firestore JSON:\n"
         "• UserModel: Identitas user & status premium keanggotaan.\n"
         "• HppModel: Parameter kalkulasi modal HPP, margin, BEP unit, BEP rupiah, nama produk, & tanggal.\n"
         "• SavingModel: Judul rencana, target nominal dana, saldo saat ini, & gambar target."),
        
        ("2. SERVICE LAYER (Data Access)", 
         "Mengelola koneksi jaringan dan database eksternal:\n"
         "• AuthService: Registrasi, login, logout, verifikasi email, & Google Sign-In.\n"
         "• FirestoreService: Operasi CRUD users, savings, finance, history, & monthly_targets.\n"
         "• PdfService: Mengenerate dokumen PDF hasil laporan keuangan."),
        
        ("3. VIEWMODEL LAYER (Business Logic)", 
         "Mengontrol state UI dan mengolah data dari layer Service:\n"
         "• AuthViewModel: State login/register & verifikasi email.\n"
         "• FinanceViewModel: Logika matematika HPP & BEP.\n"
         "• SavingViewModel: Mengelola penambahan tabungan, deposit, & tarik.\n"
         "• HomeViewModel: Target profit vs realisasi bulanan.")
    ]
    
    for i, (title, desc) in enumerate(layers):
        left = Inches(0.8 + i * 3.95)
        top = Inches(1.9)
        width = Inches(3.75)
        height = Inches(4.8)
        
        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = PRIMARY_PURPLE
        card.line.width = Pt(1.5)
        
        tf_card = card.text_frame
        tf_card.word_wrap = True
        tf_card.margin_left = Inches(0.25)
        tf_card.margin_top = Inches(0.25)
        tf_card.margin_right = Inches(0.2)
        tf_card.margin_bottom = Inches(0.2)
        
        p_title = tf_card.paragraphs[0]
        p_title.text = title
        p_title.font.name = 'Montserrat'
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = ACCENT_YELLOW
        p_title.space_after = Pt(15)
        
        p_desc = tf_card.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = 'Inter'
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = WHITE
        p_desc.line_spacing = 1.3

    # ==========================================
    # SLIDE 10: ARSITEKTUR TEKNOLOGI & SKEMA DATABASE
    # ==========================================
    slide10 = add_slide_with_header("Arsitektur Teknologi & Skema Database NoSQL")
    
    # Left Box: Tech Stack
    left_card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = PRIMARY_PURPLE
    left_card.line.width = Pt(1.5)
    
    ltf = left_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.3)
    ltf.margin_top = Inches(0.3)
    
    lp = ltf.paragraphs[0]
    lp.text = "Infrastruktur Teknologi Aplikasi"
    lp.font.name = 'Montserrat'
    lp.font.size = Pt(18)
    lp.font.bold = True
    lp.font.color.rgb = ACCENT_YELLOW
    lp.space_after = Pt(15)
    
    techs = [
        ("Flutter SDK (Dart)", "Menghasilkan UI native yang responsif dan performa tinggi (60fps) untuk Android & iOS menggunakan single codebase."),
        ("Firebase Authentication", "Sistem autentikasi aman dengan integrasi Google Sign-In dan link verifikasi email otomatis."),
        ("Cloud Firestore NoSQL", "Penyimpanan data real-time, memungkinkan data tersinkronisasi instan saat ada pembaruan saldo tabungan atau riwayat baru.")
    ]
    
    for title, desc in techs:
        p_title = ltf.add_paragraph()
        p_title.text = f"✔ {title}"
        p_title.font.bold = True
        p_title.font.size = Pt(13)
        p_title.font.color.rgb = WHITE
        
        p_desc = ltf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(10)
        
    # Right Box: Database Collections
    right_card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.9))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = CARD_BG
    right_card.line.color.rgb = PRIMARY_PURPLE
    right_card.line.width = Pt(1.5)
    
    rtf = right_card.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.3)
    rtf.margin_top = Inches(0.3)
    
    rp = rtf.paragraphs[0]
    rp.text = "Struktur Koleksi Firestore (NoSQL)"
    rp.font.name = 'Montserrat'
    rp.font.size = Pt(18)
    rp.font.bold = True
    rp.font.color.rgb = ACCENT_YELLOW
    rp.space_after = Pt(15)
    
    collections = [
        ("Koleksi 'users'", "Dokumen: UID. Fields: displayName, email, photoUrl, isPremium."),
        ("Koleksi 'finance' & 'history'", "Dokumen: Auto-generated. Fields: userId, namaProduk, totalHpp, bepUnit, bepRupiah, biayaProduksi, biayaTenagaKerja, biayaOverhead, biayaTetap, hargaJualUnit, createdAt, catatan."),
        ("Koleksi 'savings'", "Dokumen: Auto-generated. Fields: userId, title, targetAmount, currentAmount, createdAt, imageUrl."),
        ("Koleksi 'monthly_targets'", "Dokumen: userId_year_month. Fields: userId, year, month, target.")
    ]
    
    for col_name, col_desc in collections:
        p_col = rtf.add_paragraph()
        p_col.text = f"📂 {col_name}"
        p_col.font.bold = True
        p_col.font.size = Pt(12)
        p_col.font.color.rgb = WHITE
        
        p_desc = rtf.add_paragraph()
        p_desc.text = col_desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(8)

    # ==========================================
    # SLIDE 11: KESIMPULAN & Q&A (PENUTUP)
    # ==========================================
    slide11 = add_slide_with_header("Kesimpulan & Sesi Q&A")
    
    # Main Box
    main_card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.8), Inches(10.333), Inches(4.8))
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = CARD_BG
    main_card.line.color.rgb = PRIMARY_PURPLE
    main_card.line.width = Pt(2)
    
    mtf = main_card.text_frame
    mtf.word_wrap = True
    mtf.margin_left = Inches(0.5)
    mtf.margin_top = Inches(0.4)
    mtf.margin_right = Inches(0.5)
    
    mp = mtf.paragraphs[0]
    mp.text = "Mengapa Finance Tracker Menjadi Solusi Terbaik?"
    mp.font.name = 'Montserrat'
    mp.font.size = Pt(22)
    mp.font.bold = True
    mp.font.color.rgb = ACCENT_YELLOW
    mp.space_after = Pt(20)
    
    conclusions = [
        "Meningkatkan Literasi Keuangan UMKM secara mandiri dengan formula perhitungan finansial terstandar (HPP & BEP).",
        "Mencegah Kerugian Usaha melalui fitur Simulasi Harga Jual dinamis sebelum produk dipasarkan ke konsumen.",
        "Mendorong Kemandirian Finansial UMKM dengan memisahkan profit bisnis ke dalam Rencana Tabungan (Savings Plan) terarah.",
        "Menawarkan Pengalaman Pengguna Premium yang modern, andal (sinkronisasi cloud Firestore), dan efisien."
    ]
    
    for conc in conclusions:
        p_conc = mtf.add_paragraph()
        p_conc.text = "✔ " + conc
        p_conc.font.name = 'Inter'
        p_conc.font.size = Pt(14)
        p_conc.font.color.rgb = WHITE
        p_conc.space_after = Pt(12)
        
    p_thankyou = mtf.add_paragraph()
    p_thankyou.text = "\nTerima Kasih! Ada Pertanyaan? (Q&A)"
    p_thankyou.alignment = PP_ALIGN.CENTER
    p_thankyou.font.name = 'Montserrat'
    p_thankyou.font.size = Pt(20)
    p_thankyou.font.bold = True
    p_thankyou.font.color.rgb = ACCENT_YELLOW
    
    # Save the presentation
    output_path = "Presentasi_Finance_Tracker.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
